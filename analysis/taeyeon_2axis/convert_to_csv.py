"""
2026 빅콘테스트 원천 데이터(zip) -> 컬럼 분리된 CSV 변환 스크립트

- SK 유동인구 : '|' 구분자, UTF-8(BOM), 월별 6개 파일 x 3개 테이블 -> 테이블별 1개 CSV로 병합
- 신한카드     : TAB 구분자, CP949, 2개 파일 -> 각각 CSV
- 헤더는 테이블 정의서(xlsx)의 '컬럼명(한글)'으로 변환, 영문 ID와의 매핑은 csv/컬럼정의.csv 로 저장
- 202512 시간대/요일 파일은 모든 행이 정확히 2번씩 들어있어 완전중복 행을 제거
- 출력 인코딩: utf-8-sig (엑셀에서 한글 깨짐 방지)

실행: python convert_to_csv.py
"""
import io
import zipfile
from pathlib import Path

import openpyxl
import pandas as pd

BASE = Path(__file__).resolve().parent
OUT = BASE / "csv"
SK_ZIP = BASE / "원본데이터" / "2026빅콘테스트_sk통신인구_데이터_레이아웃.zip"
SH_ZIP = BASE / "원본데이터" / "2026빅콘테스트_신한카드_데이터_레이아웃.zip"


def zip_members(zpath):
    """zip 내부 파일명(cp949로 저장된 한글 파일명 복원) -> ZipInfo"""
    zf = zipfile.ZipFile(zpath)
    out = {}
    for info in zf.infolist():
        name = info.filename
        if not info.flag_bits & 0x800:  # UTF-8 플래그가 없으면 cp949로 저장된 이름
            try:
                name = name.encode("cp437").decode("cp949")
            except UnicodeError:
                pass
        out[name] = info
    return zf, out


# ---------------------------------------------------------------- 테이블 정의서 파싱
def parse_sk_definition(xlsx_bytes):
    """유동인구_테이블 정의서 -> {테이블ID: {'name','desc','cols':[(id,한글명,타입,길이,비고)]}}"""
    ws = openpyxl.load_workbook(io.BytesIO(xlsx_bytes), data_only=True).active
    tables, cur = {}, None
    for row in ws.iter_rows(values_only=True):
        if row[0] == "테이블ID":
            cur = row[2]
            tables[cur] = {"name": row[5], "desc": None, "cols": []}
        elif row[0] == "테이블설명" and cur:
            tables[cur]["desc"] = row[2]
        elif cur and isinstance(row[0], int):
            tables[cur]["cols"].append((row[1], row[2], row[3], row[4], row[8]))
    return tables


def parse_shinhan_definition(xlsx_bytes):
    """신한카드 레이아웃 -> {시트명: {'desc':[상세요건], 'cols':[(id,한글명,타입,설명)]}}"""
    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes), data_only=True)
    sheets = {}
    for ws in wb.worksheets:
        info = {"desc": [], "cols": []}
        for row in ws.iter_rows(values_only=True):
            if isinstance(row[1], int):
                info["cols"].append((row[2], row[3], row[4], row[6]))
            elif row[0] and str(row[0]).strip().startswith("-"):
                info["desc"].append(str(row[0]).strip("- ").strip())
        sheets[ws.title] = info
    return sheets


# ---------------------------------------------------------------- 변환
def convert_sk(zf, members, defs, dictionary):
    prefix_to_table = {
        "flow_age_pop": ("DMP_FLOW_POP_AGE_PCELL_GB", "SK_유동인구_성연령.csv"),
        "flow_time_pop": ("DMP_FLOW_POP_TIME_PCELL_GB", "SK_유동인구_시간대.csv"),
        "flow_wkdy_pop": ("DMP_FLOW_POP_WKDY_PCELL_GB", "SK_유동인구_요일.csv"),
    }
    for prefix, (table_id, out_name) in prefix_to_table.items():
        cols = defs[table_id]["cols"]
        rename = {c[0]: c[1] for c in cols}
        frames = []
        for name in sorted(n for n in members if n.startswith(prefix)):
            with zf.open(members[name]) as f:
                df = pd.read_csv(f, sep="|", encoding="utf-8-sig",
                                 dtype={"STD_YM": str, "BLOCK_CD": str})
            assert list(df.columns) == [c[0] for c in cols], f"{name}: 컬럼이 정의서와 다름"
            n_raw = len(df)
            df = df.drop_duplicates()
            print(f"  {name:28s} {n_raw:>9,}행 -> 중복제거 후 {len(df):>9,}행")
            frames.append(df)
        df = pd.concat(frames, ignore_index=True).rename(columns=rename)
        df.to_csv(OUT / out_name, index=False, encoding="utf-8-sig")
        print(f"  => {out_name}: {df.shape[0]:,}행 x {df.shape[1]}열")
        for cid, kname, typ, length, note in cols:
            dictionary.append({"파일": out_name, "원본테이블": table_id, "컬럼ID": cid,
                               "컬럼명": kname, "타입": typ, "길이": length, "비고": note or ""})


def convert_shinhan(zf, members, defs, dictionary):
    jobs = [
        ("신한카드_빅콘테스트2026_데이터1.txt", "주제1_데이터셋2", "신한카드_데이터1_시간대별.csv"),
        ("신한카드_빅콘테스트2026_데이터2.txt", "주제2_데이터셋2", "신한카드_데이터2_고객거주지별.csv"),
    ]
    for src, sheet, out_name in jobs:
        cols = defs[sheet]["cols"]
        with zf.open(members[src]) as f:
            df = pd.read_csv(f, sep="\t", encoding="cp949", dtype=str, keep_default_na=False)
        assert list(df.columns) == [c[0] for c in cols], f"{src}: 컬럼이 정의서와 다름"
        for c in ("TS_AT", "USE_CNT"):  # 정의서상 VARCHAR이지만 모두 정수 -> 숫자로 저장
            df[c] = pd.to_numeric(df[c]).astype("int64")
        df = df.rename(columns={c[0]: c[1] for c in cols})
        df.to_csv(OUT / out_name, index=False, encoding="utf-8-sig")
        print(f"  {src} -> {out_name}: {df.shape[0]:,}행 x {df.shape[1]}열")
        for cid, kname, typ, note in cols:
            dictionary.append({"파일": out_name, "원본테이블": sheet, "컬럼ID": cid,
                               "컬럼명": kname, "타입": typ, "길이": "", "비고": note or ""})


def main():
    OUT.mkdir(exist_ok=True)
    dictionary = []

    print("[SK 유동인구]")
    zf, members = zip_members(SK_ZIP)
    sk_def = parse_sk_definition(zf.read(members["유동인구_테이블 정의서_샘플.xlsx"]))
    convert_sk(zf, members, sk_def, dictionary)

    print("[신한카드]")
    zf, members = zip_members(SH_ZIP)
    layout = next(n for n in members if n.endswith(".xlsx"))
    sh_def = parse_shinhan_definition(zf.read(members[layout]))
    convert_shinhan(zf, members, sh_def, dictionary)

    pd.DataFrame(dictionary).to_csv(OUT / "컬럼정의.csv", index=False, encoding="utf-8-sig")
    print("  => 컬럼정의.csv 저장")


if __name__ == "__main__":
    main()
